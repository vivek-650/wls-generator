"""python-docx based DOCX text + layout extraction.

DOCX has no page/position concept comparable to a PDF, so we synthesize a
`font_size` proxy from paragraph style (Heading N / Title) and run-level bold
formatting so the same header-detection heuristics used for PDFs still work.
"""
from __future__ import annotations

from typing import Iterator, List, Union

import docx
from docx.document import Document as _DocxDocument
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .models import ExtractedTable, Line

_BASE_SIZE = 10.0
_HEADING_SIZES = {
    "title": 20.0,
    "heading 1": 16.0,
    "heading 2": 14.0,
    "heading 3": 12.0,
}


def _paragraph_font_size(paragraph) -> float:
    style_name = (paragraph.style.name or "").strip().lower() if paragraph.style else ""
    if style_name in _HEADING_SIZES:
        return _HEADING_SIZES[style_name]

    # Explicit run-level font size, if the document author set one.
    sizes = [
        run.font.size.pt
        for run in paragraph.runs
        if run.font is not None and run.font.size is not None and run.text.strip()
    ]
    if sizes:
        return max(sizes)
    return _BASE_SIZE


def _paragraph_is_bold(paragraph) -> bool:
    runs_with_text = [r for r in paragraph.runs if r.text.strip()]
    if not runs_with_text:
        return False
    bold_chars = sum(len(r.text) for r in runs_with_text if r.bold)
    total_chars = sum(len(r.text) for r in runs_with_text)
    return total_chars > 0 and (bold_chars / total_chars) >= 0.5


def _iter_block_items(parent) -> Iterator[Union[Paragraph, Table]]:
    """Yield each top-level paragraph and table in *true document order*.

    `document.paragraphs` and `document.tables` are separate flat
    collections in python-docx — iterating them one after the other (as a
    naive implementation would) silently moves every table's content to
    the very end of the document regardless of where it actually sits,
    corrupting section boundaries for any resume that puts a table (e.g. a
    Skills table) anywhere but the last section. Walking the underlying
    XML body directly and reconstructing each `w:p`/`w:tbl` child as its
    proper wrapper object is the standard fix.
    """
    parent_elm = parent.element.body if isinstance(parent, _DocxDocument) else parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _iter_cell_block_items(cell: _Cell) -> Iterator[Union[Paragraph, Table]]:
    yield from _iter_block_items(cell)


def extract_lines(file_bytes: bytes) -> List[Line]:
    """Parse a DOCX byte string into an ordered list of `Line` objects, in
    true document order (paragraphs and tables interleaved as they
    actually appear — see `_iter_block_items`).

    Raises on malformed DOCX bytes (the caller translates that into a 422).
    """
    import io

    document = docx.Document(io.BytesIO(file_bytes))
    lines: List[Line] = []
    order = 0

    def add_paragraph(paragraph: Paragraph) -> None:
        nonlocal order
        text = paragraph.text.strip()
        if not text:
            return
        lines.append(
            Line(
                text=text,
                page=1,
                order=order,
                font_size=_paragraph_font_size(paragraph),
                bold=_paragraph_is_bold(paragraph),
            )
        )
        order += 1

    for block in _iter_block_items(document):
        if isinstance(block, Paragraph):
            add_paragraph(block)
        else:
            for row in block.rows:
                for cell in row.cells:
                    for cell_block in _iter_cell_block_items(cell):
                        if isinstance(cell_block, Paragraph):
                            add_paragraph(cell_block)
                        # Nested tables inside a cell are rare in resumes;
                        # not worth the added complexity of recursing here.

    return lines


def extract_tables(file_bytes: bytes) -> List[ExtractedTable]:
    """DOCX tables are a first-class structure in the file format itself —
    `document.tables` gives unambiguous row/column cell text directly, no
    heuristics needed (unlike PDF, where a "table" is just visually
    positioned text unless the source drew ruling lines).
    """
    import io

    document = docx.Document(io.BytesIO(file_bytes))
    tables: List[ExtractedTable] = []
    for table in document.tables:
        rows = [[" ".join(cell.text.split()).strip() for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(row)]
        if rows:
            tables.append(ExtractedTable(page=1, rows=rows))
    return tables
